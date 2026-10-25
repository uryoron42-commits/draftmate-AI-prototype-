import re
from io import BytesIO

from docx import Document
from docx.shared import Pt


def _runs(par, line):
    for part in re.split(r"(\*\*.+?\*\*)", line):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            par.add_run(part[2:-2]).bold = True
        elif part:
            par.add_run(part)


def generate_docx(title, content, verification=None):
    d = Document()
    d.styles["Normal"].font.name = "Times New Roman"
    d.styles["Normal"].font.size = Pt(12)
    d.add_heading(title, 0)
    for line in content.splitlines():
        s = line.strip()
        if not s:
            continue
        if re.match(r"^#{1,4}\s", s):
            d.add_heading(s.lstrip("#").strip(), min(3, len(s) - len(s.lstrip("#"))))
        elif re.match(r"^[-*•]\s+", s):
            _runs(d.add_paragraph(style="List Bullet"), re.sub(r"^[-*•]\s+", "", s))
        else:
            _runs(d.add_paragraph(), s)
    if verification and verification.get("citations"):
        d.add_page_break()
        d.add_heading("Verification Certificate", 1)
        t = d.add_table(rows=1, cols=3)
        t.style = "Light Grid Accent 1"
        for i, h in enumerate(("Authority", "Status", "Notes")):
            t.rows[0].cells[i].text = h
        for r in verification["citations"]:
            c = t.add_row().cells
            c[0].text, c[1].text = r["raw"], r["status"].capitalize()
            c[2].text = (r.get("note", "") + (f" Suggested: {r['replacement']}" if r.get("replacement") else "")).strip()
        d.add_paragraph("\nAI-assisted draft. Reviewed by: ______________________   Date: ____________")
    b = BytesIO(); d.save(b)
    return b.getvalue()
